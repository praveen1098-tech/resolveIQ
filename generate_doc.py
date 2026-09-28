import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# Set standard margins
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Colors
PRIMARY = RGBColor(15, 23, 42)    # Slate 900
ACCENT = RGBColor(2, 132, 199)    # Sky 600
MUTED = RGBColor(100, 116, 139)   # Slate 500

def set_run_font(run, font_name="Calibri", size_pt=11, color=PRIMARY, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic

# Title
title_p = doc.add_paragraph()
title_run = title_p.add_run("ResolveIQ: SRE Incident Response & Memory Agent")
set_run_font(title_run, size_pt=24, color=PRIMARY, bold=True)
title_p.paragraph_format.space_after = Pt(4)

# Subtitle
sub_p = doc.add_paragraph()
sub_run = sub_p.add_run("Technical Overview, Memory Architecture & Operational Workflow")
set_run_font(sub_run, size_pt=13, color=MUTED, italic=True)
sub_p.paragraph_format.space_after = Pt(20)

# 1. Executive Summary
h1 = doc.add_heading("1. Executive Summary", level=1)
h1.style.font.color.rgb = PRIMARY
p1 = doc.add_paragraph(
    "ResolveIQ is an autonomous Site Reliability Engineering (SRE) incident triage agent "
    "powered by Hindsight persistent memory and high-throughput Groq LLM inference. "
    "Traditional diagnostic agents act in a stateless manner, treating every recurring outage "
    "as an unprecedented failure. ResolveIQ retains technical post-mortems, verified root causes, "
    "and runbook remediations across deployment cycles, closing the feedback loop on microservice uptime."
)
p1.paragraph_format.space_after = Pt(12)

# 2. Key Architecture Elements
h2 = doc.add_heading("2. System Architecture & Component Mapping", level=1)
h2.style.font.color.rgb = PRIMARY

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Component"
hdr_cells[1].text = "Technology"
hdr_cells[2].text = "Functional Role"

for cell in hdr_cells:
    for p in cell.paragraphs:
        for r in p.runs:
            set_run_font(r, size_pt=10, color=PRIMARY, bold=True)

data = [
    ("Memory Layer", "Hindsight Cloud (Vectorize)", "Semantic retrieval of past incidents (recall) and storage of validated fixes (retain)."),
    ("Reasoning Engine", "Groq (Llama / Qwen)", "Synthesizes live telemetry with historical incident context to isolate root cause."),
    ("Triage Console", "Streamlit", "Dual-pane interface providing live alert inputs and inspection of retrieved memories."),
    ("Data Store", "JSON Schemas", "Standardized structures for incident metadata, services, and runbook definitions.")
]

for item, tech, role in data:
    row_cells = table.add_row().cells
    row_cells[0].text = item
    row_cells[1].text = tech
    row_cells[2].text = role
    for cell in row_cells:
        for p in cell.paragraphs:
            for r in p.runs:
                set_run_font(r, size_pt=9.5, color=PRIMARY)

doc.add_paragraph().paragraph_format.space_after = Pt(14)

# 3. Memory Lifecycle
h3 = doc.add_heading("3. The Hindsight Continuous Learning Loop", level=1)
h3.style.font.color.rgb = PRIMARY

bp1 = doc.add_paragraph(style='List Bullet')
r1 = bp1.add_run("Recall Stage: ")
set_run_font(r1, bold=True)
bp1.add_run("When an alert fires (e.g., HTTP 503 on auth-service), the agent issues a semantic query to the 'resolveiq' Hindsight bank, retrieving related past incidents, error patterns, and successful resolutions.")

bp2 = doc.add_paragraph(style='List Bullet')
r2 = bp2.add_run("Reasoning & Triage: ")
set_run_font(r2, bold=True)
bp2.add_run("The Groq LLM receives both the live error logs and the retrieved memories, allowing it to prescribe the proven historical resolution (e.g., patching a missing Kubernetes ConfigMap key) rather than general troubleshooting.")

bp3 = doc.add_paragraph(style='List Bullet')
r3 = bp3.add_run("Retain Stage: ")
set_run_font(r3, bold=True)
bp3.add_run("Once an engineer confirms the actual fix, the post-mortem is retained back into Hindsight with metadata tags, ensuring the organization never repeats the diagnostic cycle for that failure mode.")

# 4. Mandatory References
h4 = doc.add_heading("4. Documentation & Repository Links", level=1)
h4.style.font.color.rgb = PRIMARY

ref_p = doc.add_paragraph()
ref_p.add_run("• Hindsight GitHub: https://github.com/vectorize-io/hindsight\n")
ref_p.add_run("• Hindsight Documentation: https://hindsight.vectorize.io/\n")
ref_p.add_run("• Vectorize Agent Memory: https://vectorize.io/what-is-agent-memory\n")

doc.save("ResolveIQ_Project_Overview.docx")
print("✅ ResolveIQ_Project_Overview.docx successfully generated!")
